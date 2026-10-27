# -*- coding: utf-8 -*-
"""
生成开票数据 Excel 模板
用法：python make_template.py
"""
from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

ROOT = Path(__file__).resolve().parent
OUT = ROOT / "data" / "开票数据.xlsx"

HEADER = [
    "序号", "客户名称", "出库单号", "出库日期", "价税合计",
    "开票人", "业务日期", "发票日期",
    "日期区间例外", "摘要", "备注", "发票号",
    "应收单号", "发票单据号", "凭证号", "处理状态", "错误信息",
]

# 前 12 列是输入，后 5 列是脚本回写
INPUT_N = 12

SAMPLE = [
    # 客户名称, 出库单号, 出库日期, 价税合计, 开票人, 业务日期, 发票日期, 区间例外, 摘要, 备注, 发票号
    [1, "示例客户A", "XSCKD000001", "2026-08-28", 11300.00, "张三", "2026-09-05", "2026-09-05", "", "示例客户A8月货款", "", ""],
    [2, "示例客户A", "XSCKD000002", "2026-09-02", 5650.00, "张三", "2026-09-05", "2026-09-05", "", "示例客户A8月货款", "", ""],
    [3, "示例客户B", "XSCKD000003", "2026-09-10", 22600.00, "张三", "2026-09-05", "2026-09-05", "2026-09-01~2026-09-30", "示例客户B9月货款", "特例：区间非默认", ""],
]

DOC = [
    ("列", "含义", "是否必填", "说明"),
    ("序号", "行号", "否", "方便对照，可不填"),
    ("客户名称", "客户", "是", "【分组依据之一】同一张票的客户必须一致"),
    ("出库单号", "销售出库单编号", "是", "脚本据此在金蝶列表里勾选，如 XSCKD000001"),
    ("出库日期", "出库单业务日期", "否", "仅作核对参考"),
    ("价税合计", "本行金额", "是", "组内会求和，与金蝶生成的发票金额核对，不一致会报警"),
    ("开票人", "填进金蝶发票的「开票人」", "是", "例如 张三"),
    ("业务日期", "填进金蝶发票的「业务日期」", "是", "例如 2026-09-05"),
    ("发票日期", "填进金蝶发票的「发票日期」", "是", "【分组依据之一】同客户同一发票日期的算一张票"),
    ("日期区间例外", "非默认的出库单日期区间", "否", "留空=按默认「上月26号~当月25号」；填了写 2026-08-26~2026-09-25 格式"),
    ("摘要", "凭证摘要", "否", "留空则用 config.yaml 的 digest_template 生成；同一张票填一样"),
    ("备注", "特例说明", "否", "给自己看的，脚本不处理"),
    ("发票号", "选填", "否", "不用填。只有当你需要打破「客户+发票日期」的默认分组时才填：填了同一发票号的行会强制归为一组"),
    ("应收单号", "【脚本回写】", "-", "下推生成的应收单编号"),
    ("发票单据号", "【脚本回写】", "-", "金蝶销售增值税专用发票单据号"),
    ("凭证号", "【脚本回写】", "-", "生成的凭证号"),
    ("处理状态", "【脚本回写】", "-", "完成 / 失败"),
    ("错误信息", "【脚本回写】", "-", "失败原因"),
]


def main():
    OUT.parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook()

    # ---------- Sheet1: 数据 ----------
    ws = wb.active
    ws.title = "开票数据"

    thin = Side(style="thin", color="BFBFBF")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    fill_in = PatternFill("solid", fgColor="EAF3FB")   # 输入列
    fill_out = PatternFill("solid", fgColor="F2F2F2")  # 回写列
    fill_hd = PatternFill("solid", fgColor="D9E2F3")

    for i, h in enumerate(HEADER, 1):
        c = ws.cell(row=1, column=i, value=h)
        c.font = Font(bold=True, size=11)
        c.fill = fill_hd if i <= INPUT_N else PatternFill("solid", fgColor="E8E8E8")
        c.border = border
        c.alignment = Alignment(horizontal="center", vertical="center")

    for r_i, row in enumerate(SAMPLE, start=2):
        for c_i, v in enumerate(row, start=1):
            c = ws.cell(row=r_i, column=c_i, value=v)
            c.border = border
            c.fill = fill_in if c_i <= INPUT_N else fill_out
        # 日期/金额列格式化：4=出库日期 5=价税合计 7=业务日期 8=发票日期
        ws.cell(row=r_i, column=4).number_format = "yyyy-mm-dd"
        ws.cell(row=r_i, column=5).number_format = "#,##0.00"
        ws.cell(row=r_i, column=7).number_format = "yyyy-mm-dd"
        ws.cell(row=r_i, column=8).number_format = "yyyy-mm-dd"

    widths = [6, 20, 16, 12, 12, 10, 12, 12, 24, 28, 22, 16,
              18, 18, 14, 10, 30]
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w

    ws.freeze_panes = "A2"
    ws.row_dimensions[1].height = 24

    # ---------- Sheet2: 填写说明 ----------
    ws2 = wb.create_sheet("填写说明")
    for r_i, row in enumerate(DOC, start=1):
        for c_i, v in enumerate(row, start=1):
            c = ws2.cell(row=r_i, column=c_i, value=v)
            c.border = border
            if r_i == 1:
                c.font = Font(bold=True)
                c.fill = fill_hd
    for i, w in enumerate([18, 26, 10, 60], 1):
        ws2.column_dimensions[get_column_letter(i)].width = w

    tips = [
        "分组规则：同一张票的多张出库单，各占一行。默认按「客户名称 + 发票日期」分为一组，"
        "下推时选「分录合并」生成一张金蝶发票。",
        "如果同客户同一天要开多张票，就在「发票号」列填上不同的号来强制拆分（该列平时留空即可）。",
        "填进金蝶的字段只有三个：开票人、业务日期、发票日期。发票号码不用填。",
        "示例行只是演示格式，正式跑之前请删掉或改成真实数据。",
    ]
    for i, t in enumerate(tips):
        c = ws2.cell(row=len(DOC) + 2 + i, column=1, value=t)
        c.font = Font(bold=(i < 3), color="C00000")

    wb.save(OUT)
    print(f"模板已生成：{OUT}")
    print("  Sheet「开票数据」：填数据")
    print("  Sheet「填写说明」：字段含义")


if __name__ == "__main__":
    main()
